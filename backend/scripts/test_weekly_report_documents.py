import json
import os
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile
import subprocess
from datetime import date
from pathlib import Path
from xml.etree import ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'src/scripts'))
import generate_weekly_summary_docx as generator

GROUPS = [
    ('personnel', 'Lĩnh vực Nội vụ và công tác cán bộ'),
    ('education', 'Giáo dục và đào tạo'),
    ('social', 'Bảo trợ xã hội và bảo hiểm y tế'),
    ('health', 'Y tế'),
    ('veterans', 'Người có công'),
    ('administration', 'Tổ chức bộ máy, công tác cán bộ và cải cách hành chính'),
    ('digital', 'Khoa học, công nghệ và chuyển đổi số'),
    ('culture', 'Văn hóa và các lĩnh vực liên quan'),
]


def fixture():
    config = json.loads((ROOT / 'src/config/document-defaults.json').read_text(encoding='utf-8'))
    return {
        'config': config, 'groups': [{'id': key, 'title': title} for key, title in GROUPS],
        'startDate': '2026-10-18T17:00:00Z', 'dueDate': '2026-10-22T16:59:59Z',
        'entries': [
            {'id': f'{key}-{section}', 'groupId': key, 'section': section,
             'text': f'Nguyên văn {title}: 692 trẻ; 536 trẻ; 77,4%; số 2324/QĐ-UBND.\n**Nội dung đậm** và <u>gạch chân</u>; 5 < 10.'}
            for key, title in GROUPS for section in ('results', 'nextTasks')
        ],
    }


class DocumentTests(unittest.TestCase):
    def test_staff_and_monthly_html_template_exports(self):
        from docx import Document
        payload = dict(period='Tháng 10/2026', reportTitle='BÁO CÁO', department='QA', sender='QA', field='Y tế',
                       startDate='2026-10-19T00:00:00Z', endDate='2026-10-22T00:00:00Z', dueDate='2026-10-22T00:00:00Z',
                       contentFormat='html-v1', content='<p><strong>Rich template</strong> **literal** &lt; 10</p><ol><li><p>First</p></li><li><p>Second</p></li></ol>',
                       nextTasks='<p><u>Next task</u></p>')
        with tempfile.TemporaryDirectory() as folder:
            for script, template in [('generate_weekly_docx.py', 'weekly_template.docx'), ('docx_summary_generator.py', 'monthly_template.docx')]:
                output = Path(folder) / f'{script}.docx'
                result = subprocess.run([sys.executable, str(ROOT / 'src/scripts' / script), str(ROOT / 'src/templates' / template), str(output)], input=json.dumps(payload).encode('utf-8'), capture_output=True, timeout=30)
                self.assertEqual(result.returncode, 0, result.stderr.decode('utf-8', errors='replace'))
                actual = Document(output)
                text = '\n'.join(p.text for p in actual.paragraphs)
                self.assertIn('**literal** < 10', text)
                self.assertTrue(any(run.bold and run.text == 'Rich template' for p in actual.paragraphs for run in p.runs))
                self.assertTrue(any(run.underline and run.text == 'Next task' for p in actual.paragraphs for run in p.runs))
                self.assertTrue(any(p._p.pPr is not None and p._p.pPr.numPr is not None for p in actual.paragraphs))
                self.assertIn(b'9000', actual.part.numbering_part.blob)

    def test_html_marks_lists_and_literals(self):
        from docx import Document
        payload = fixture()
        payload['entries'] = [{ 'id': 'rich', 'groupId': 'health', 'section': 'results', 'contentFormat': 'html-v1',
            'text': '<p><strong>Bold <em>both</em></strong> <u>Underline</u> **literal** &lt; 10 &amp; 20</p><ol start="3"><li><p>First</p><ul><li><p>Nested</p></li></ul></li><li><p>Second</p></li></ol><p>End<br>Next line</p>' }]
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / 'report.docx'
            generator.build_document(payload, output)
            document = Document(output)
            text = '\n'.join(p.text for p in document.paragraphs)
            self.assertIn('**literal** < 10 & 20', text)
            self.assertIn('End\nNext line', text)
            self.assertTrue(any(run.bold and run.italic and run.text == 'both' for p in document.paragraphs for run in p.runs))
            self.assertTrue(any(run.underline and run.text == 'Underline' for p in document.paragraphs for run in p.runs))
            ns = {'w': generator.qn('w:p').split('}')[0][1:]}
            with zipfile.ZipFile(output) as package:
                root = ET.fromstring(package.read('word/document.xml'))
                self.assertEqual(len(root.findall('.//w:numPr', ns)), 3)
                numbering = ET.fromstring(package.read('word/numbering.xml'))
                self.assertTrue(any(node.get(generator.qn('w:val')) == '3' for node in numbering.findall('.//w:start', ns)))
            if os.environ.get('RUN_PDF_TESTS') == '1':
                import pymupdf
                pdf = output.with_suffix('.pdf')
                generator.convert_pdf(output, pdf)
                with pymupdf.open(pdf) as actual:
                    extracted = '\n'.join(page.get_text() for page in actual)
                self.assertIn('**literal** < 10 & 20', extracted)
                self.assertRegex(extracted, r'3\.\s+First')
                self.assertRegex(extracted, r'4\.\s+Second')
                self.assertIn('Nested', extracted)

    def test_year_boundaries(self):
        for value, label in [('2026-10-19', 'tuần 43 năm 2026'), ('2027-01-01', 'tuần 53 năm 2026'), ('2027-01-04', 'tuần 1 năm 2027'), ('2024-12-30', 'tuần 1 năm 2025')]:
            self.assertEqual(generator.week_label(date.fromisoformat(value)), label)
        self.assertEqual(generator.local_date('2026-10-18T17:00:00Z'), date(2026, 10, 19))

    def test_word_contents_and_format(self):
        payload = fixture()
        payload['config']['signerName'] = 'Người ký kiểm thử'
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / 'report.docx'
            generator.build_document(payload, output)
            with zipfile.ZipFile(output) as package:
                xml = package.read('word/document.xml')
                root = ET.fromstring(xml)
                ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
                text = '\n'.join(node.text or '' for node in root.findall('.//w:t', ns))
                self.assertIn('Kết quả công tác tuần 43 năm 2026', text)
                self.assertIn('và phương hướng, nhiệm vụ tuần 44 năm 2026', text)
                self.assertIn('Số:       /BC-VHXH', text)
                self.assertIn('I. KẾT QUẢ THỰC HIỆN NHIỆM VỤ TUẦN 43\n', text)
                self.assertIn('II. PHƯƠNG HƯỚNG, NHIỆM VỤ TUẦN 44\n', text)
                self.assertIn('cập nhật đến ngày 22/10/2026', text)
                self.assertIn('Trên đây là báo cáo kết quả công tác tuần 43 và phương hướng, nhiệm vụ tuần 44 năm 2026', text)
                self.assertIn('kính báo cáo UBND xã xem xét, chỉ đạo./.', text)
                self.assertIn('- Lưu: VHXH.', text)
                self.assertNotIn('____________', text)
                self.assertEqual(len(root.findall('.//w:pBdr/w:bottom', ns)), 3)
                self.assertIn('Người ký kiểm thử', text)
                self.assertIn('5 < 10.', text)
                self.assertEqual(text.count('692 trẻ; 536 trẻ; 77,4%; số 2324/QĐ-UBND.'), 16)
                for index, (_, title) in enumerate(GROUPS, 1):
                    self.assertEqual(text.count(f'{index}. {title}'), 2)
                self.assertNotIn('III.', text)
                self.assertEqual(len(root.findall('.//w:tbl', ns)), 2)
                self.assertTrue(root.findall('.//w:rPr/w:b', ns))
                self.assertTrue(root.findall('.//w:rPr/w:u', ns))
                self.assertFalse(any(name.startswith('word/media/') for name in package.namelist()))
                self.assertIn('Times New Roman', xml.decode())
                today = generator.datetime.now(generator.VIETNAM_TZ).date()
                self.assertIn(f'ngày {today.day:02d} tháng {today.month:02d} năm {today.year}', text)
                self.assertIn('PAGE', package.read('word/header1.xml').decode())

    def test_cutoff_uses_period_due_date_not_generation_time(self):
        payload = fixture()
        payload['sourceGeneratedAt'] = '2026-10-30T10:00:00Z'
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / 'report.docx'
            generator.build_document(payload, output)
            from docx import Document
            text = '\n'.join(p.text for p in Document(output).paragraphs)
            self.assertIn('cập nhật đến ngày 22/10/2026', text)
            self.assertNotIn('cập nhật đến ngày 30/10/2026', text)

    def test_unknown_group_fails(self):
        payload = fixture()
        payload['entries'][0]['groupId'] = None
        with tempfile.TemporaryDirectory() as folder:
            with self.assertRaisesRegex(ValueError, 'chưa phân nhóm'):
                generator.build_document(payload, Path(folder) / 'bad.docx')

    def test_missing_pdf_converter(self):
        with patch.dict(os.environ, {}, clear=True), patch.object(generator.shutil, 'which', return_value=None), patch.object(generator.os, 'name', 'posix'):
            with self.assertRaisesRegex(RuntimeError, 'Chưa có LibreOffice'):
                generator.convert_pdf(Path('missing.docx'), Path('missing.pdf'))

    @unittest.skipUnless(os.environ.get('RUN_PDF_TESTS') == '1', 'Set RUN_PDF_TESTS=1 with LibreOffice or Windows Word')
    def test_real_pdf_and_render(self):
        import pymupdf
        folder = Path(os.environ.get('WEEKLY_QA_DIR', tempfile.gettempdir())) / 'weekly-document-test'
        folder.mkdir(exist_ok=True)
        source, target = folder / 'report.docx', folder / 'report.pdf'
        generator.build_document(fixture(), source)
        generator.convert_pdf(source, target)
        with pymupdf.open(target) as document:
            text = ''.join(page.get_text() for page in document)
            self.assertIn('tuần 43 năm 2026', text)
            self.assertIn('tuần 44 năm 2026', text)
            self.assertIn('Lê Việt Thanh', text)
            self.assertEqual(text.count('77,4%'), 16)
            for index, page in enumerate(document, 1):
                page.get_pixmap(matrix=pymupdf.Matrix(1.2, 1.2)).save(str(folder / f'page-{index}.png'))
        print(f'PDF render for visual QA: {folder}')


if __name__ == '__main__':
    unittest.main()
