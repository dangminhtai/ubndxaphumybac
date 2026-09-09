#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generate a native Microsoft Excel (.xlsx) file for Weekly Tasks ("Nhiệm vụ tuần")
following the official agency template and spreadsheet design standards.
"""

import sys
import json
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


def create_weekly_tasks_excel(payload, output_path):
    week_label = payload.get("weekLabel", "Tuần").strip()
    tasks = payload.get("tasks", [])

    wb = Workbook()
    ws = wb.active
    ws.title = "Nhiệm vụ tuần"

    # Ensure grid lines are visible
    ws.views.sheetView[0].showGridLines = True

    # Typography & Styles
    FONT_NAME = "Arial"

    title_font = Font(name=FONT_NAME, size=14, bold=True, color="000000")
    header_font = Font(name=FONT_NAME, size=11, bold=True, color="000000")
    name_font = Font(name=FONT_NAME, size=11, bold=True, color="000000")
    body_font = Font(name=FONT_NAME, size=11, color="000000")

    # Alignments
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    stt_align = Alignment(horizontal="center", vertical="top")
    name_align = Alignment(horizontal="left", vertical="top", wrap_text=True)
    content_align = Alignment(horizontal="left", vertical="top", wrap_text=True)
    deadline_align = Alignment(horizontal="center", vertical="top", wrap_text=True)

    # Borders
    thin_border_side = Side(style="thin", color="000000")
    cell_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side,
    )

    # Header Fill (Very subtle professional fill or clean white)
    header_fill = PatternFill(start_color="F2F4F8", end_color="F2F4F8", fill_type="solid")

    # Column Widths
    col_widths = {
        "A": 8,    # STT
        "B": 28,   # HỌ VÀ TÊN
        "C": 75,   # NHIỆM VỤ TUẦN
        "D": 25,   # THỜI GIAN HOÀN THÀNH
    }
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    # Row 1: Spacing
    ws.row_dimensions[1].height = 8

    # Row 2: Title
    ws.row_dimensions[2].height = 32
    title_text = f"NHIỆM VỤ {week_label.upper()}"
    ws.merge_cells("A2:D2")
    title_cell = ws["A2"]
    title_cell.value = title_text
    title_cell.font = title_font
    title_cell.alignment = center_align

    # Row 3: Blank spacer
    ws.row_dimensions[3].height = 12

    # Row 4: Table Headers
    ws.row_dimensions[4].height = 28
    headers = [
        ("A4", "STT"),
        ("B4", "HỌ VÀ TÊN"),
        ("C4", "NHIỆM VỤ TUẦN"),
        ("D4", "THỜI GIAN HOÀN THÀNH"),
    ]
    for cell_id, text in headers:
        cell = ws[cell_id]
        cell.value = text
        cell.font = header_font
        cell.alignment = header_align
        cell.fill = header_fill
        cell.border = cell_border

    # Rows 5+: Data rows
    start_row = 5
    for idx, item in enumerate(tasks, start=1):
        row_num = start_row + idx - 1

        user_info = item.get("userId", {})
        if isinstance(user_info, dict):
            full_name = user_info.get("fullName", "").strip()
        else:
            full_name = str(user_info or "").strip()

        content = (item.get("content") or "").strip()
        deadline = (item.get("deadline") or "").strip()

        # Cell A: STT
        c_stt = ws.cell(row=row_num, column=1, value=idx)
        c_stt.font = body_font
        c_stt.alignment = stt_align
        c_stt.border = cell_border

        # Cell B: HỌ VÀ TÊN (Uppercase & Bold)
        c_name = ws.cell(row=row_num, column=2, value=full_name.upper())
        c_name.font = name_font
        c_name.alignment = name_align
        c_name.border = cell_border

        # Cell C: NHIỆM VỤ TUẦN (Multiline text with line breaks preserved)
        c_content = ws.cell(row=row_num, column=3, value=content)
        c_content.font = body_font
        c_content.alignment = content_align
        c_content.border = cell_border

        # Cell D: THỜI GIAN HOÀN THÀNH
        c_deadline = ws.cell(row=row_num, column=4, value=deadline)
        c_deadline.font = body_font
        c_deadline.alignment = deadline_align
        c_deadline.border = cell_border

        # Set appropriate row height based on content lines
        if content:
            line_count = len(content.split("\n"))
            # Estimate additional lines from text wrapping (approx 65 chars per line in 75 width col)
            wrap_extra = sum(max(0, len(line) // 65) for line in content.split("\n"))
            total_lines = max(1, line_count + wrap_extra)
            # ~18pt per line + padding
            ws.row_dimensions[row_num].height = max(24, total_lines * 19 + 6)
        else:
            ws.row_dimensions[row_num].height = 24

    # Page Print Setup
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "4:4"  # Repeat header row on multiple pages

    # Ensure target directory exists
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

    wb.save(output_path)
    return output_path


def main():
    if len(sys.argv) < 3:
        print("Usage: generate_weekly_tasks_excel.py <input_json_path> <output_xlsx_path>", file=sys.stderr)
        sys.exit(1)

    input_path = sys.argv[1]
    output_path = sys.argv[2]

    try:
        with open(input_path, "r", encoding="utf-8") as f:
            payload = json.load(f)

        create_weekly_tasks_excel(payload, output_path)
        print(f"SUCCESS: Generated {output_path}")
    except Exception as e:
        print(f"ERROR: {str(e)}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
