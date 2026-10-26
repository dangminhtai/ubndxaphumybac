"""Weekly department report; keep the shared staff/monthly generator unchanged."""
import json
import os
import shutil
import subprocess
import sys
import tempfile
from datetime import datetime, timedelta
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx_summary_generator import VIETNAM_TZ, markdown_to_html, parse_html_to_lines, parse_runs


def local_date(value):
    return datetime.fromisoformat(value.replace('Z', '+00:00')).astimezone(VIETNAM_TZ).date()


def week_label(day):
    year, week, _ = day.isocalendar()
    return f'tuần {week} năm {year}'


def paragraph(container, text='', *, bold=False, italic=False, align=None, size=14, indent=False, after=6, keep=False):
    p = container.add_paragraph()
    p.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    f = p.paragraph_format
    f.space_after, f.space_before = Pt(after), Pt(0)
    f.line_spacing = 1.05
    f.first_line_indent = Cm(1.25) if indent else Cm(0)
    f.keep_with_next, f.widow_control = keep, True
    for value, b, i, u in parse_runs(markdown_to_html(text)):
        r = p.add_run(value)
        r.bold, r.italic, r.underline = b or bold, i or italic, u
        r.font.name, r.font.size = 'Times New Roman', Pt(size)
        r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    return p


def body_text(doc, text):
    for line in parse_html_to_lines(markdown_to_html(text)):
        paragraph(doc, line, indent=True)


def two_columns(doc, left_width, right_width):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width, table.columns[1].width = Cm(left_width), Cm(right_width)
    for cell, width in zip(table.rows[0].cells, (left_width, right_width)):
        cell.width = Cm(width)
        cell._element.remove(cell.paragraphs[0]._element)
        margins = OxmlElement('w:tcMar')
        for side in ('top', 'left', 'bottom', 'right'):
            item = OxmlElement(f'w:{side}')
            item.set(qn('w:w'), '0')
            item.set(qn('w:type'), 'dxa')
            margins.append(item)
        cell._tc.get_or_add_tcPr().append(margins)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    return table.rows[0].cells


def build_document(payload, output):
    config, groups, entries = payload['config'], payload['groups'], payload['entries']
    valid_ids = {group['id'] for group in groups}
    if any(entry['text'].strip() and entry.get('groupId') not in valid_ids for entry in entries):
        raise ValueError('Còn nội dung chưa phân nhóm')
    start = local_date(payload['startDate'])
    current, following = week_label(start), week_label(start + timedelta(days=7))
    today = datetime.now(VIETNAM_TZ).date()
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2), Cm(2)
    section.left_margin, section.right_margin = Cm(3), Cm(1.5)
    section.header_distance = Cm(0.8)
    style = doc.styles['Normal']
    style.font.name, style.font.size = 'Times New Roman', Pt(14)
    section.different_first_page_header_footer = True
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    header._p.append(field)

    left, right = two_columns(doc, 6.5, 10)
    center = WD_ALIGN_PARAGRAPH.CENTER
    paragraph(left, config['parentAgency'], align=center, size=13, after=0)
    paragraph(left, config['agency'], bold=True, align=center, size=13, after=0)
    paragraph(left, '____________', align=center, size=10, after=6)
    paragraph(left, f"Số:       /{config['documentCode']}", align=center, size=13)
    paragraph(right, 'CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM', bold=True, align=center, size=13, after=0)
    paragraph(right, 'Độc lập - Tự do - Hạnh phúc', bold=True, align=center, size=14, after=0)
    paragraph(right, '________________________', align=center, size=10, after=6)
    paragraph(right, f"{config['place']}, ngày {today.day:02d} tháng {today.month:02d} năm {today.year}", italic=True, align=center, size=13)
    paragraph(doc, 'BÁO CÁO', bold=True, align=center, after=0, keep=True)
    paragraph(doc, f'Kết quả công tác {current}', bold=True, align=center, after=0, keep=True)
    paragraph(doc, f'và phương hướng, nhiệm vụ {following}', bold=True, align=center, after=14, keep=True)
    paragraph(doc, f"Kính gửi: {config['addressee']}.", align=center, after=12)
    cutoff = f", cập nhật đến ngày {local_date(payload['sourceGeneratedAt']).strftime('%d/%m/%Y')}" if payload.get('sourceGeneratedAt') else ''
    paragraph(doc, f"{config['agencyDisplayName']} báo cáo kết quả thực hiện nhiệm vụ {current}{cutoff}, và phương hướng, nhiệm vụ {following} như sau:", indent=True)
    for key, title in [('results', f'I. KẾT QUẢ THỰC HIỆN NHIỆM VỤ {current.upper()}'), ('nextTasks', f'II. PHƯƠNG HƯỚNG, NHIỆM VỤ {following.upper()}')]:
        paragraph(doc, title, bold=True, keep=True, indent=True)
        for index, group in enumerate(groups, 1):
            group_entries = [entry for entry in entries if entry['section'] == key and entry['groupId'] == group['id'] and entry['text'].strip()]
            paragraph(doc, f"{index}. {group['title']}", bold=True, keep=bool(group_entries), indent=True)
            for entry in group_entries:
                body_text(doc, entry['text'])
    paragraph(doc, f"Trên đây là báo cáo kết quả công tác {current} và phương hướng, nhiệm vụ {following} của {config['agencyDisplayName']}, kính báo cáo {config['addressee']} xem xét, chỉ đạo./.", indent=True, keep=True)
    left, right = two_columns(doc, 8, 8.5)
    paragraph(left, 'Nơi nhận:', bold=True, italic=True, size=12, after=0)
    for recipient in config['recipients']:
        paragraph(left, f'- {recipient};', size=11, after=0)
    paragraph(right, config['signerTitle'], bold=True, align=center, after=60, keep=True)
    paragraph(right, config['signerName'], bold=True, align=center)
    doc.save(output)


def convert_pdf(source, target):
    soffice = os.environ.get('SOFFICE_BIN') or shutil.which('libreoffice') or shutil.which('soffice')
    if not soffice and os.name == 'nt':
        candidate = Path(os.environ.get('PROGRAMFILES', 'C:/Program Files')) / 'LibreOffice/program/soffice.exe'
        if candidate.exists():
            soffice = str(candidate)
    if soffice:
        with tempfile.TemporaryDirectory(prefix='weekly-lo-') as profile:
            result = subprocess.run([soffice, f'-env:UserInstallation={Path(profile).as_uri()}', '--headless', '--convert-to', 'pdf:writer_pdf_Export', '--outdir', str(target.parent), str(source)], capture_output=True, timeout=90)
            if result.returncode or not target.exists():
                raise RuntimeError('LibreOffice không tạo được PDF: ' + result.stderr.decode('utf-8', errors='replace'))
    elif os.name == 'nt':
        import pythoncom
        import win32com.client
        pythoncom.CoInitialize()
        word = document = None
        try:
            word = win32com.client.DispatchEx('Word.Application')
            word.Visible, word.DisplayAlerts = False, 0
            document = word.Documents.Open(str(source.resolve()), ReadOnly=True, AddToRecentFiles=False)
            document.ExportAsFixedFormat(str(target.resolve()), 17)
        finally:
            if document is not None:
                document.Close(False)
            if word is not None:
                word.Quit()
            pythoncom.CoUninitialize()
    else:
        raise RuntimeError('Chưa có LibreOffice để xuất PDF. Cài LibreOffice và font Times New Roman, hoặc đặt SOFFICE_BIN.')
    if not target.exists() or target.read_bytes()[:5] != b'%PDF-':
        raise RuntimeError('Không tạo được file PDF hợp lệ')


def main():
    if len(sys.argv) != 2:
        raise ValueError('Usage: generate_weekly_summary_docx.py <output.docx|output.pdf>')
    output = Path(sys.argv[1]).resolve()
    if output.suffix not in ('.docx', '.pdf'):
        raise ValueError('Định dạng không hỗ trợ')
    payload = json.loads(sys.stdin.buffer.read().decode('utf-8-sig'))
    output.parent.mkdir(parents=True, exist_ok=True)
    source = output.with_suffix('.docx')
    try:
        build_document(payload, source)
        if output.suffix == '.pdf':
            convert_pdf(source, output)
    finally:
        if output.suffix == '.pdf':
            source.unlink(missing_ok=True)


if __name__ == '__main__':
    main()
